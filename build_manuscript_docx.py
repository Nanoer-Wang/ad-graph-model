"""Editable Word text and high-resolution display equations from shared source."""
from pathlib import Path
from html.parser import HTMLParser
import re
from docx import Document
from docx.shared import Cm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from md_to_html_paper import SRC,parse,inline
from build_pdf import image_size,rows_from_html,col_widths,TEXT_WIDTH

ROOT=Path(__file__).resolve().parent
OUT=ROOT/'paper_v3/Manuscript_Revised.docx'

class Runs(HTMLParser):
    def __init__(self,p):super().__init__();self.p=p;self.stack=[]
    def handle_starttag(self,tag,attrs):self.stack.append(tag)
    def handle_endtag(self,tag):
        if tag in self.stack:self.stack.remove(tag)
    def handle_data(self,s):
        r=self.p.add_run(s);r.bold=any(t in self.stack for t in ['b','strong'])
        r.italic=any(t in self.stack for t in ['i','em'])
        r.font.subscript='sub' in self.stack;r.font.superscript='sup' in self.stack

def runs(p,s):Runs(p).feed(s)
def main():
    doc=Document();sec=doc.sections[0]
    sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.5)
    for st in ['Normal','Title','Heading 1','Heading 2','Caption']:
        style=doc.styles[st];style.font.name='Cambria';style.font.color.rgb=RGBColor(0,0,0)
        style.paragraph_format.widow_control=True
    normal=doc.styles['Normal'];normal.font.size=Pt(11)
    normal.paragraph_format.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing=1.2;normal.paragraph_format.space_after=Pt(6)
    doc.styles['Title'].font.size=Pt(17)
    doc.styles['Heading 1'].font.size=Pt(13);doc.styles['Heading 2'].font.size=Pt(11.5)
    doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False
    for st in ['Heading 1','Heading 2']:
        doc.styles[st].paragraph_format.alignment=WD_ALIGN_PARAGRAPH.LEFT
    title,cover,blocks=parse(SRC.read_text(encoding='utf-8'))
    p=doc.add_paragraph(title,'Title');p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for text in cover:
        p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;runs(p,inline(text))
    doc.add_page_break()
    foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    img=None;eqnum=0
    for kind,payload in blocks:
        if kind in ['h2','h3']:
            p=doc.add_paragraph(style='Heading 1' if kind=='h2' else 'Heading 2');runs(p,payload)
        elif kind=='figimg':img=re.search('src="images/([^"]+)"',payload).group(1)
        elif kind=='figcap':
            if img:
                p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
                w,h=image_size(img);p.add_run().add_picture(str(SRC.parent/'images'/img),width=Cm(w),height=Cm(h));img=None
            p=doc.add_paragraph(style='Caption');runs(p,payload)
        elif kind=='eq':
            name=re.search('src="images/([^"]+)"',payload).group(1);eqnum+=1;w,h=image_size(name)
            t=doc.add_table(rows=1,cols=2);t.autofit=False;t.columns[0].width=Cm(15);t.columns[1].width=Cm(1)
            p=t.cell(0,0).paragraphs[0];p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.add_run().add_picture(str(SRC.parent/'images'/name),width=Cm(w),height=Cm(h))
            t.cell(0,1).paragraphs[0].add_run(f'({eqnum})')
            for cell in t.rows[0].cells:cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            doc.add_paragraph().paragraph_format.space_after=Pt(0)
        elif kind=='table':
            caption,rows=rows_from_html(payload)
            p=doc.add_paragraph(style='Caption');p.paragraph_format.keep_with_next=True;runs(p,caption)
            t=doc.add_table(rows=len(rows),cols=len(rows[0]));t.autofit=False
            widths=col_widths(rows,TEXT_WIDTH)
            for j,ww in enumerate(widths):t.columns[j].width=Cm(ww/TEXT_WIDTH*16)
            for i,row in enumerate(rows):
                for j,val in enumerate(row):
                    cell=t.cell(i,j);cell.width=Cm(widths[j]/TEXT_WIDTH*16)
                    p=cell.paragraphs[0];p.paragraph_format.space_after=Pt(4);p.paragraph_format.space_before=Pt(4)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                    runs(p,val)
                    for r in p.runs:r.font.size=Pt(9);r.bold=i==0
                    if i==0:
                        sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEEEEE');cell._tc.get_or_add_tcPr().append(sh)
                trpr=t.rows[i]._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
            borders=OxmlElement('w:tblBorders')
            for side in ['top','left','bottom','right','insideH','insideV']:
                el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
            t._tbl.tblPr.append(borders)
        elif kind=='refs':
            for i,text in enumerate(re.findall('<li class="reference-item">(.*?)</li>',payload,re.S),1):
                p=doc.add_paragraph();p.paragraph_format.left_indent=Cm(.8);p.paragraph_format.first_line_indent=Cm(-.8)
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                runs(p,f'[{i}] '+text)
                for r in p.runs:r.font.size=Pt(9.5)
        else:
            p=doc.add_paragraph();runs(p,inline(payload) if kind=='p' else payload)
    OUT.parent.mkdir(exist_ok=True);doc.core_properties.title=title;doc.core_properties.author=''
    doc.save(OUT);print(f'DOCX saved: {len(doc.paragraphs)} paragraphs, {len(doc.tables)} tables, {len(doc.inline_shapes)} images')

if __name__=='__main__':main()
