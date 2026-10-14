"""Stage 3: convert the typeset manuscript HTML into the final .docx.

Thin wrapper around the `html-to-docx` skill converter.  The process working
directory is switched to the HTML folder because html4docx resolves relative
``<img src>`` paths against the CWD.
"""
from __future__ import annotations

import os
import sys
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parent
STAGE2 = ROOT / "output" / "ad_passivation_paper" / "stage2"
HTML = STAGE2 / "formatted-fibril-interface-passivation.html"
OUT = ROOT / "paper_v2" / "Manuscript_Fibril_Interface_Passivation_AD_Model.docx"

PLUGIN = (Path.home() / ".workbuddy" / "plugins" / "cache" / "workbuddy-builtin"
          / "tencent-docx")
SKILL_SCRIPTS = sorted(PLUGIN.glob("*/skills/html-to-docx/scripts"))
if not SKILL_SCRIPTS:
    raise SystemExit("html-to-docx skill scripts not found")
sys.path.insert(0, str(SKILL_SCRIPTS[-1]))

from html_to_docx import convert, ConvertOptions  # noqa: E402


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)
    os.chdir(STAGE2)  # relative image paths resolve from here

    html_text = HTML.read_text(encoding="utf-8")
    result = convert(
        html_text,
        output_path=str(OUT),
        options=ConvertOptions(page_size="A4", margin_top=2.5, margin_bottom=2.5,
                               margin_left=2.5, margin_right=2.5,
                               base_dir=str(STAGE2)),
    )
    if not result.success:
        raise SystemExit(f"conversion failed: {result.error}")

    with zipfile.ZipFile(OUT) as z:
        media = [n for n in z.namelist() if "/media/" in n]
    from docx import Document  # noqa: E402
    doc = Document(str(OUT))
    missing = [p.text for p in doc.paragraphs if p.text.startswith("<image:")]

    print(f"docx:        {OUT}")
    print(f"paragraphs:  {len(doc.paragraphs)}")
    print(f"tables:      {len(doc.tables)}")
    print(f"images:      {len(media)}")
    print(f"unresolved:  {len(missing)} {missing[:3]}")
    print(f"warnings:    {result.warnings[:5]}")


if __name__ == "__main__":
    main()
